"""stats API 集成测试 - dashboard 聚合查询、days 参数、缓存、排名"""
import pytest
from datetime import datetime, timezone, timedelta
from db.models import Workflow, RunRecord
from routers.stats import _dashboard_cache


@pytest.fixture(autouse=True)
def clear_stats_cache():
    """每个测试前清空 dashboard 缓存,避免跨测试污染"""
    _dashboard_cache.clear()
    yield
    _dashboard_cache.clear()


@pytest.mark.asyncio
async def test_dashboard_empty_db(client):
    """空库返回零值不报错"""
    resp = await client.get("/api/stats/dashboard?days=7")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total_runs"] == 0
    assert data["success_count"] == 0
    assert data["failed_count"] == 0
    assert data["success_rate"] == 0.0
    assert data["avg_time_ms"] == 0
    assert len(data["trend"]) == 7
    assert data["workflow_ranking"] == []
    assert data["failure_clusters"] == []


@pytest.mark.asyncio
async def test_dashboard_with_runs(client, test_db):
    """有执行记录返回正确聚合"""
    # 插入 1 个工作流 + 2 个执行记录
    wf = Workflow(id="wf-1", name="测试工作流", steps="[]", edges="[]")
    test_db.add(wf)
    test_db.commit()

    now = datetime.now(timezone.utc)
    for i, status in enumerate(["success", "failed"]):
        run = RunRecord(
            id=f"run-{i}",
            workflow_id="wf-1",
            trigger_type="manual",
            status=status,
            total_time_ms=1000 + i * 500,
            started_at=now - timedelta(hours=i),
            finished_at=now - timedelta(hours=i) + timedelta(seconds=1),
            error="连接超时" if status == "failed" else None,
        )
        test_db.add(run)
    test_db.commit()

    resp = await client.get("/api/stats/dashboard?days=7")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total_runs"] == 2
    assert data["success_count"] == 1
    assert data["failed_count"] == 1
    assert data["success_rate"] == 50.0


@pytest.mark.asyncio
async def test_dashboard_days_param(client, test_db):
    """days=7/30/90 返回对应天数趋势"""
    for days in [7, 30, 90]:
        _dashboard_cache.clear()
        resp = await client.get(f"/api/stats/dashboard?days={days}")
        assert resp.status_code == 200
        data = resp.json()
        assert len(data["trend"]) == days, f"days={days} 时趋势长度应为 {days}"


@pytest.mark.asyncio
async def test_dashboard_days_validation(client):
    """days 超范围(0/91)返回 422"""
    resp0 = await client.get("/api/stats/dashboard?days=0")
    assert resp0.status_code == 422

    resp91 = await client.get("/api/stats/dashboard?days=91")
    assert resp91.status_code == 422


@pytest.mark.asyncio
async def test_dashboard_cache_hit(client, test_db):
    """连续调用第二次命中缓存(返回相同数据)"""
    wf = Workflow(id="wf-c", name="缓存测试", steps="[]", edges="[]")
    test_db.add(wf)
    test_db.commit()
    run = RunRecord(
        id="run-c1",
        workflow_id="wf-c",
        trigger_type="manual",
        status="success",
        total_time_ms=500,
        started_at=datetime.now(timezone.utc),
        finished_at=datetime.now(timezone.utc),
    )
    test_db.add(run)
    test_db.commit()

    resp1 = await client.get("/api/stats/dashboard?days=7")
    assert resp1.status_code == 200
    data1 = resp1.json()
    assert data1["total_runs"] == 1

    # 第二次调用应命中缓存(返回相同数据)
    resp2 = await client.get("/api/stats/dashboard?days=7")
    assert resp2.status_code == 200
    data2 = resp2.json()
    assert data2["total_runs"] == 1
    assert data2 == data1


@pytest.mark.asyncio
async def test_workflow_ranking(client, test_db):
    """按执行数降序排列"""
    # 插入 2 个工作流,A 执行 3 次,B 执行 1 次
    for wid, name in [("wf-a", "工作流A"), ("wf-b", "工作流B")]:
        test_db.add(Workflow(id=wid, name=name, steps="[]", edges="[]"))
    test_db.commit()

    now = datetime.now(timezone.utc)
    # wf-a: 3 次执行
    for i in range(3):
        test_db.add(RunRecord(
            id=f"run-a-{i}",
            workflow_id="wf-a",
            trigger_type="manual",
            status="success",
            total_time_ms=100,
            started_at=now - timedelta(hours=i),
            finished_at=now,
        ))
    # wf-b: 1 次执行
    test_db.add(RunRecord(
        id="run-b-0",
        workflow_id="wf-b",
        trigger_type="manual",
        status="success",
        total_time_ms=200,
        started_at=now,
        finished_at=now,
    ))
    test_db.commit()

    resp = await client.get("/api/stats/dashboard?days=7")
    assert resp.status_code == 200
    data = resp.json()
    ranking = data["workflow_ranking"]
    assert len(ranking) >= 2
    # wf-a 应排第一(执行数 3 > 1)
    assert ranking[0]["workflow_id"] == "wf-a"
    assert ranking[0]["total"] == 3
    assert ranking[1]["workflow_id"] == "wf-b"
    assert ranking[1]["total"] == 1


# ========== D2: Dashboard 统计报表导出 ==========

@pytest.mark.asyncio
async def test_export_stats_xlsx_empty_db(client):
    """空库导出 xlsx:返回有效 Excel,含全部基础 sheet(无 Token sheet)"""
    resp = await client.get("/api/stats/export?format=xlsx")
    assert resp.status_code == 200
    assert "spreadsheet" in resp.headers["content-type"]
    assert resp.content[:2] == b"PK"
    from openpyxl import load_workbook
    import io as _io
    wb = load_workbook(_io.BytesIO(resp.content))
    # 基础 sheet 必须存在
    for sheet_name in ["总览", "趋势", "工具排名", "工作流排名", "失败聚类", "耗时分位数"]:
        assert sheet_name in wb.sheetnames, f"缺少 sheet: {sheet_name}"
    # 空 db 无 token 数据,Token用量 sheet 不应存在
    assert "Token用量" not in wb.sheetnames
    # 总览 sheet 第一行数据应为 0
    ws = wb["总览"]
    assert ws.cell(2, 1).value == "总执行数"
    assert ws.cell(2, 2).value == 0


@pytest.mark.asyncio
async def test_export_stats_xlsx_with_data(client, test_db):
    """有执行数据时导出 xlsx:总览含非零值,趋势 sheet 含数据行"""
    wf = Workflow(id="wf-x1", name="导出测试流", steps="[]", edges="[]")
    test_db.add(wf)
    test_db.commit()
    now = datetime.now(timezone.utc)
    for i, status in enumerate(["success", "failed"]):
        test_db.add(RunRecord(
            id=f"run-x-{i}",
            workflow_id="wf-x1",
            trigger_type="manual",
            status=status,
            total_time_ms=1000 + i * 500,
            started_at=now,
            finished_at=now,
            error="连接超时" if status == "failed" else None,
        ))
    test_db.commit()

    resp = await client.get("/api/stats/export?format=xlsx&days=7")
    assert resp.status_code == 200
    from openpyxl import load_workbook
    import io as _io
    wb = load_workbook(_io.BytesIO(resp.content))
    # 总览:总执行数 = 2
    ws = wb["总览"]
    assert ws.cell(2, 2).value == 2
    # 工作流排名 sheet 含导出测试流
    ws_wf = wb["工作流排名"]
    assert ws_wf.cell(2, 2).value == "导出测试流"
    assert ws_wf.cell(2, 3).value == 2  # 执行数
    # 失败聚类 sheet 含"网络超时"(error="连接超时"匹配)
    ws_fc = wb["失败聚类"]
    assert ws_fc.cell(2, 1).value == "网络超时"


@pytest.mark.asyncio
async def test_export_stats_unsupported_format(client):
    """不支持的格式返回 400"""
    resp = await client.get("/api/stats/export?format=csv")
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_export_stats_default_format_is_xlsx(client):
    """默认 format=xlsx"""
    resp = await client.get("/api/stats/export")
    assert resp.status_code == 200
    assert "spreadsheet" in resp.headers["content-type"]


@pytest.mark.asyncio
async def test_export_stats_days_param(client, test_db):
    """days=30 时趋势 sheet 有 30 行数据"""
    resp = await client.get("/api/stats/export?format=xlsx&days=30")
    assert resp.status_code == 200
    from openpyxl import load_workbook
    import io as _io
    wb = load_workbook(_io.BytesIO(resp.content))
    ws = wb["趋势"]
    # 1 表头 + 30 天数据
    assert ws.max_row == 31


@pytest.mark.asyncio
async def test_export_stats_filename_chinese(client):
    """文件名含中文,通过 RFC 5987 编码"""
    resp = await client.get("/api/stats/export?format=xlsx&days=7")
    assert resp.status_code == 200
    cd = resp.headers["content-disposition"]
    assert "filename*=UTF-8''" in cd
    # "FlowGenie报表_7天.xlsx" 的百分比编码应含 %E6%8A%A5 (报)
    assert "%E6%8A%A5" in cd


# ========== P2: 索引存在性 + 缓存命中验证 ==========

def test_run_records_status_started_at_index_exists(test_engine):
    """P2: 验证 run_records 表的复合索引 (status, started_at) 已创建。

    该索引加速 dashboard 中 WHERE status IN (...) AND started_at >= ? 类查询。
    通过查询 sqlite_master 确认索引存在,避免 ORM 隐式建索引失败被忽略。
    """
    with test_engine.connect() as conn:
        rows = conn.exec_driver_sql(
            "SELECT name FROM sqlite_master "
            "WHERE type='index' AND tbl_name='run_records'"
        ).fetchall()
        index_names = {r[0] for r in rows}
    assert "ix_run_records_status_started_at" in index_names, (
        f"复合索引缺失,实际索引: {index_names}"
    )
    # 同时验证 workflow_id 单列索引仍存在(基础索引不应被破坏)
    assert "ix_run_records_workflow_id" in index_names


def test_workflow_versions_index_exists(test_engine):
    """P2: 验证 workflow_versions 表的索引存在(回归保护)。"""
    with test_engine.connect() as conn:
        rows = conn.exec_driver_sql(
            "SELECT name FROM sqlite_master "
            "WHERE type='index' AND tbl_name='workflow_versions'"
        ).fetchall()
        index_names = {r[0] for r in rows}
    assert "ix_workflow_versions_workflow_id" in index_names


@pytest.mark.asyncio
async def test_dashboard_cache_invalidation_on_workflow_change(client, test_db):
    """P2: 工作流写操作后 dashboard 缓存应被失效,避免读到陈旧数据。

    虽然 workflows 路由与 stats 路由的缓存独立,但 dashboard 数据依赖
    run_records,本测试验证 invalidate_dashboard_cache 清空缓存后,新增
    run_record 能被立即聚合(而非命中旧缓存)。
    """
    from routers.stats import _dashboard_cache, invalidate_dashboard_cache

    # 第一次请求,缓存空,触发聚合
    _dashboard_cache.clear()
    resp1 = await client.get("/api/stats/dashboard?days=7")
    assert resp1.status_code == 200
    assert resp1.json()["total_runs"] == 0
    # 缓存应被填充
    assert len(_dashboard_cache) > 0

    # 插入新 run_record 后主动失效缓存(模拟执行完成后的调用)
    wf = Workflow(id="wf-inv", name="失效测试", steps="[]", edges="[]")
    test_db.add(wf)
    test_db.commit()
    run = RunRecord(
        id="run-inv",
        workflow_id="wf-inv",
        trigger_type="manual",
        status="success",
        total_time_ms=300,
        started_at=datetime.now(timezone.utc),
        finished_at=datetime.now(timezone.utc),
    )
    test_db.add(run)
    test_db.commit()

    invalidate_dashboard_cache()
    assert len(_dashboard_cache) == 0

    # 再次请求应重新聚合,看到新增的记录
    resp2 = await client.get("/api/stats/dashboard?days=7")
    assert resp2.status_code == 200
    assert resp2.json()["total_runs"] == 1


@pytest.mark.asyncio
async def test_dashboard_cache_keyed_by_days(client):
    """P2: 不同 days 参数使用不同缓存条目,互不干扰。"""
    from routers.stats import _dashboard_cache

    _dashboard_cache.clear()
    # 请求 days=7 与 days=30,各自独立缓存
    await client.get("/api/stats/dashboard?days=7")
    await client.get("/api/stats/dashboard?days=30")
    assert "dashboard:7" in _dashboard_cache
    assert "dashboard:30" in _dashboard_cache
    assert len(_dashboard_cache) >= 2


def test_tool_timeout_env_default(monkeypatch):
    """P2: TOOL_TIMEOUT 环境变量未设置时,默认值为 120s。"""
    # 清除环境变量,重新导入模块验证默认值
    monkeypatch.delenv("TOOL_TIMEOUT", raising=False)
    import importlib
    import engine.executor as executor_mod
    importlib.reload(executor_mod)
    assert executor_mod.DEFAULT_TOOL_TIMEOUT == 120


def test_tool_timeout_env_override(monkeypatch):
    """P2: 设置 TOOL_TIMEOUT=60 时,默认超时变为 60s。"""
    monkeypatch.setenv("TOOL_TIMEOUT", "60")
    import importlib
    import engine.executor as executor_mod
    importlib.reload(executor_mod)
    assert executor_mod.DEFAULT_TOOL_TIMEOUT == 60


def test_tool_timeout_env_clamped_to_max(monkeypatch):
    """P2: TOOL_TIMEOUT 超过 300s 上限时被钳制为 300s(防止误配置)。"""
    monkeypatch.setenv("TOOL_TIMEOUT", "9999")
    import importlib
    import engine.executor as executor_mod
    importlib.reload(executor_mod)
    assert executor_mod.DEFAULT_TOOL_TIMEOUT == 300


def test_tool_timeout_env_invalid_falls_back(monkeypatch):
    """P2: TOOL_TIMEOUT 非法值(非数字)回退到默认 120s。"""
    monkeypatch.setenv("TOOL_TIMEOUT", "not-a-number")
    import importlib
    import engine.executor as executor_mod
    importlib.reload(executor_mod)
    assert executor_mod.DEFAULT_TOOL_TIMEOUT == 120


# 恢复 executor 模块到默认状态,避免影响后续测试
@pytest.fixture(autouse=True)
def _restore_executor_module(monkeypatch):
    """每个测试后恢复 TOOL_TIMEOUT 环境与 executor 模块默认值"""
    yield
    monkeypatch.delenv("TOOL_TIMEOUT", raising=False)
    import importlib
    import engine.executor as executor_mod
    importlib.reload(executor_mod)
